"""
Gerador de documentos DOCX e PDF profissionais.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


class DocumentGenerator:
    """Gera memoriais descritivos em DOCX/PDF com tipografia e estrutura profissionais."""

    def __init__(self) -> None:
        self.document: Document | None = None

    def generate_docx(self, content: Dict[str, Any], metadata: Dict[str, Any]) -> io.BytesIO:
        self.document = Document()
        self._setup_page()
        self._setup_styles()
        self._add_cover_page(metadata)

        for index, section in enumerate(content.get("sections", []), start=1):
            self._add_section(section, index)

        if content.get("calculations"):
            self._add_calculations(content["calculations"])

        if content.get("normas"):
            self._add_norms(content["normas"])

        self._add_project_data_block(metadata)
        self._add_page_numbers()

        stream = io.BytesIO()
        self.document.save(stream)
        stream.seek(0)
        return stream

    def _setup_page(self) -> None:
        assert self.document is not None
        for section in self.document.sections:
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(3.0)
            section.right_margin = Cm(2.0)
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)

    def _setup_styles(self) -> None:
        assert self.document is not None
        styles = self.document.styles

        normal = styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_after = Pt(8)
        pf.first_line_indent = Cm(1.25)

        for level, size in ((1, 14), (2, 12), (3, 12)):
            style = styles[f"Heading {level}"]
            style.font.name = "Times New Roman"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            style.paragraph_format.space_after = Pt(8)
            style.paragraph_format.first_line_indent = Cm(0)

    def _add_page_numbers(self) -> None:
        assert self.document is not None
        for section in self.document.sections:
            footer = section.footer
            footer.is_linked_to_previous = False
            paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.clear()
            run = paragraph.add_run("Página ")
            run.font.size = Pt(10)
            run.font.name = "Times New Roman"

            fld_char_begin = OxmlElement("w:fldChar")
            fld_char_begin.set(qn("w:fldCharType"), "begin")
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = " PAGE "
            fld_char_end = OxmlElement("w:fldChar")
            fld_char_end.set(qn("w:fldCharType"), "end")

            run2 = paragraph.add_run()
            run2._r.append(fld_char_begin)
            run2._r.append(instr)
            run2._r.append(fld_char_end)
            run2.font.size = Pt(10)

            run3 = paragraph.add_run(" de ")
            run3.font.size = Pt(10)

            fld2_begin = OxmlElement("w:fldChar")
            fld2_begin.set(qn("w:fldCharType"), "begin")
            instr2 = OxmlElement("w:instrText")
            instr2.set(qn("xml:space"), "preserve")
            instr2.text = " NUMPAGES "
            fld2_end = OxmlElement("w:fldChar")
            fld2_end.set(qn("w:fldCharType"), "end")
            run4 = paragraph.add_run()
            run4._r.append(fld2_begin)
            run4._r.append(instr2)
            run4._r.append(fld2_end)
            run4.font.size = Pt(10)

    @staticmethod
    def _display(value: Any, fallback: str = "Não informado") -> str:
        if value is None or value == "":
            return fallback
        return str(value)

    def _add_cover_page(self, metadata: Dict[str, Any]) -> None:
        assert self.document is not None
        professional = metadata.get("professional", {}) or {}

        for _ in range(2):
            self.document.add_paragraph()

        title = self.document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.first_line_indent = Cm(0)
        run = title.add_run(self._display(metadata.get("title"), "MEMORIAL DESCRITIVO").upper())
        run.bold = True
        run.font.size = Pt(18)
        run.font.name = "Times New Roman"

        if metadata.get("subtitle"):
            sub = self.document.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub.paragraph_format.first_line_indent = Cm(0)
            r = sub.add_run(self._display(metadata.get("subtitle")))
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(64, 64, 64)

        self.document.add_paragraph()
        self.document.add_paragraph()

        for label, value in (
            ("Obra", metadata.get("obra")),
            ("Local", metadata.get("local")),
            ("Responsável técnico", metadata.get("responsavel") or professional.get("full_name")),
            ("CREA", metadata.get("crea")),
            ("Empresa", professional.get("company_name")),
            ("Data", datetime.now().strftime("%d/%m/%Y")),
        ):
            line = self.document.add_paragraph()
            line.alignment = WD_ALIGN_PARAGRAPH.CENTER
            line.paragraph_format.first_line_indent = Cm(0)
            line.paragraph_format.space_after = Pt(4)
            a = line.add_run(f"{label}: ")
            a.bold = True
            a.font.size = Pt(12)
            b = line.add_run(self._display(value))
            b.font.size = Pt(12)

        self.document.add_paragraph()
        note = self.document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.first_line_indent = Cm(0)
        nr = note.add_run(
            "Documento técnico elaborado conforme normas da ABNT aplicáveis à disciplina."
        )
        nr.italic = True
        nr.font.size = Pt(10)

        self.document.add_page_break()

    def _write_body_blocks(self, text: str) -> None:
        """Quebra texto em parágrafos e listas a partir de quebras de linha."""
        assert self.document is not None
        blocks = re.split(r"\n\s*\n", (text or "").strip())
        if not blocks:
            blocks = [""]

        for block in blocks:
            lines = [line.strip() for line in block.splitlines() if line.strip()]
            if not lines:
                continue

            bullet_lines = [
                line
                for line in lines
                if line.startswith(("- ", "• ", "* "))
                or re.match(r"^\d+[\.\)]\s+", line)
            ]
            if len(bullet_lines) == len(lines) and len(lines) > 1:
                for line in lines:
                    clean = re.sub(r"^([-•*]|\d+[\.\)])\s+", "", line)
                    p = self.document.add_paragraph(clean, style="List Bullet")
                    p.paragraph_format.first_line_indent = Cm(0)
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)
                continue

            paragraph = self.document.add_paragraph(" ".join(lines))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

    def _add_section(self, section: Dict[str, Any], index: int) -> None:
        assert self.document is not None
        title = self._display(section.get("title"), f"{index}. SEÇÃO")
        heading = self.document.add_heading(title, level=1)
        heading.paragraph_format.first_line_indent = Cm(0)
        self._write_body_blocks(self._display(section.get("content"), ""))

        for subsection in section.get("subsections") or []:
            sub_title = self._display(subsection.get("title"), "Subseção")
            sub_heading = self.document.add_heading(sub_title, level=2)
            sub_heading.paragraph_format.first_line_indent = Cm(0)
            self._write_body_blocks(self._display(subsection.get("content"), ""))

    def _add_calculations(self, calculations: Dict[str, Any]) -> None:
        assert self.document is not None
        self.document.add_page_break()
        heading = self.document.add_heading("CÁLCULOS E DIMENSIONAMENTOS", level=1)
        heading.paragraph_format.first_line_indent = Cm(0)

        intro = self.document.add_paragraph(
            "Os valores a seguir foram obtidos a partir das fórmulas parametrizadas do "
            "modelo técnico e dos dados de entrada do empreendimento. Devem ser "
            "conferidos pelo responsável técnico antes da emissão definitiva."
        )
        intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for calc_name, calc_data in calculations.items():
            if "error" in calc_data:
                continue
            p = self.document.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.add_run(f"{calc_name}: ").bold = True
            value = calc_data.get("value")
            unit = calc_data.get("unit", "")
            if isinstance(value, (int, float)):
                p.add_run(f"{value:.2f} {unit}".strip())
            else:
                p.add_run(f"{value} {unit}".strip())
            formula = self.document.add_paragraph()
            formula.paragraph_format.first_line_indent = Cm(0)
            fr = formula.add_run(f"Fórmula: {calc_data.get('formula', '')}")
            fr.italic = True
            fr.font.size = Pt(10)

    def _add_norms(self, normas: list) -> None:
        assert self.document is not None
        self.document.add_page_break()
        heading = self.document.add_heading("REFERÊNCIAS NORMATIVAS", level=1)
        heading.paragraph_format.first_line_indent = Cm(0)
        intro = self.document.add_paragraph(
            "O desenvolvimento deste memorial observa, no que couber, as seguintes "
            "normas técnicas e documentos de referência, além da legislação local "
            "aplicável e das instruções do corpo de bombeiros da unidade federativa:"
        )
        intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for norma in normas:
            p = self.document.add_paragraph(str(norma), style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)

    def _add_project_data_block(self, metadata: Dict[str, Any]) -> None:
        assert self.document is not None
        professional = metadata.get("professional", {}) or {}
        shared = metadata.get("shared_data", {}) or {}

        self.document.add_page_break()
        heading = self.document.add_heading("DADOS DO PROJETO E RESPONSABILIDADE TÉCNICA", level=1)
        heading.paragraph_format.first_line_indent = Cm(0)

        info_table = self.document.add_table(rows=11, cols=2)
        info_table.style = "Table Grid"
        info_data = [
            ("Obra", self._display(metadata.get("obra"))),
            ("Local", self._display(metadata.get("local") or shared.get("localizacao"))),
            ("Município", self._display(shared.get("municipio"))),
            ("UF", self._display(shared.get("uf"))),
            ("Área construída (m²)", self._display(shared.get("area_construida"))),
            ("Pavimentos", self._display(shared.get("numero_pavimentos"))),
            ("Responsável técnico", self._display(metadata.get("responsavel") or professional.get("full_name"))),
            ("Título profissional", self._display(professional.get("professional_title"))),
            ("CREA", self._display(metadata.get("crea"))),
            ("Empresa", self._display(professional.get("company_name"))),
            ("ART/RRT", self._display(professional.get("art_number") or professional.get("rrt_number"))),
        ]
        for i, (label, value) in enumerate(info_data):
            info_table.rows[i].cells[0].text = label
            info_table.rows[i].cells[1].text = value
            for cell in info_table.rows[i].cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(11)

        self.document.add_paragraph()
        closing = self.document.add_paragraph(
            "Declaro, para os devidos fins, que as informações técnicas constantes deste "
            "memorial foram elaboradas sob minha responsabilidade profissional, em "
            "conformidade com as normas técnicas vigentes e com as premissas de projeto "
            "informadas pelo contratante. Qualquer alteração de uso, carga, geometria ou "
            "sistema construtivo deverá ser submetida a nova análise técnica."
        )
        closing.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        self.document.add_paragraph()
        signature = self.document.add_paragraph()
        signature.paragraph_format.first_line_indent = Cm(0)
        signature.add_run("\n\n____________________________________________\n")
        signature.add_run(
            f"{self._display(professional.get('full_name') or metadata.get('responsavel'), 'Responsável técnico')}\n"
        )
        signature.add_run(
            f"{self._display(professional.get('professional_title'), 'Engenheiro(a)')} — "
            f"CREA {self._display(metadata.get('crea'))}\n"
        )
        signature.add_run(f"Data: {datetime.now().strftime('%d/%m/%Y')}")

    def generate_pdf(self, content: Dict[str, Any], metadata: Dict[str, Any]) -> io.BytesIO:
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            ListFlowable,
            ListItem,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
        from reportlab.lib import colors

        def display(value: Any, fallback: str = "Não informado") -> str:
            return str(value) if value not in (None, "") else fallback

        def esc(text: str) -> str:
            return (
                display(text, "")
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace("\n", "<br/>")
            )

        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=3 * cm,
            rightMargin=2 * cm,
            topMargin=2.5 * cm,
            bottomMargin=2.5 * cm,
            title=display(metadata.get("title"), "Memorial Descritivo"),
        )
        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle(
                name="CoverTitle",
                parent=styles["Title"],
                fontName="Times-Bold",
                fontSize=16,
                alignment=TA_CENTER,
                spaceAfter=14,
            )
        )
        styles.add(
            ParagraphStyle(
                name="BodyJust",
                parent=styles["BodyText"],
                fontName="Times-Roman",
                fontSize=11,
                leading=16,
                alignment=TA_JUSTIFY,
                spaceAfter=10,
                firstLineIndent=20,
            )
        )
        styles.add(
            ParagraphStyle(
                name="MetaCenter",
                parent=styles["Normal"],
                fontName="Times-Roman",
                fontSize=11,
                leading=14,
                alignment=TA_CENTER,
                spaceAfter=4,
            )
        )
        styles.add(
            ParagraphStyle(
                name="H1Pt",
                parent=styles["Heading1"],
                fontName="Times-Bold",
                fontSize=13,
                spaceBefore=12,
                spaceAfter=8,
            )
        )
        styles.add(
            ParagraphStyle(
                name="H2Pt",
                parent=styles["Heading2"],
                fontName="Times-Bold",
                fontSize=12,
                spaceBefore=10,
                spaceAfter=6,
            )
        )

        story: List[Any] = []
        professional = metadata.get("professional", {}) or {}
        shared = metadata.get("shared_data", {}) or {}

        story.append(Spacer(1, 1.5 * cm))
        story.append(Paragraph(esc(display(metadata.get("title"), "MEMORIAL DESCRITIVO").upper()), styles["CoverTitle"]))
        if metadata.get("subtitle"):
            story.append(Paragraph(esc(metadata.get("subtitle")), styles["MetaCenter"]))
        story.append(Spacer(1, 1.2 * cm))
        for label, value in (
            ("Obra", metadata.get("obra")),
            ("Local", metadata.get("local")),
            ("Responsável técnico", metadata.get("responsavel") or professional.get("full_name")),
            ("CREA", metadata.get("crea")),
            ("Empresa", professional.get("company_name")),
            ("Data", datetime.now().strftime("%d/%m/%Y")),
        ):
            story.append(Paragraph(f"<b>{label}:</b> {esc(display(value))}", styles["MetaCenter"]))
        story.append(Spacer(1, 0.8 * cm))
        story.append(
            Paragraph(
                "<i>Documento técnico elaborado conforme normas da ABNT aplicáveis à disciplina.</i>",
                styles["MetaCenter"],
            )
        )
        story.append(PageBreak())

        for section in content.get("sections", []):
            story.append(Paragraph(esc(section.get("title") or "Seção"), styles["H1Pt"]))
            for block in re.split(r"\n\s*\n", display(section.get("content"), "").strip()) or [""]:
                story.append(Paragraph(esc(block), styles["BodyJust"]))
            for subsection in section.get("subsections") or []:
                story.append(Paragraph(esc(subsection.get("title") or "Subseção"), styles["H2Pt"]))
                for block in re.split(r"\n\s*\n", display(subsection.get("content"), "").strip()) or [""]:
                    story.append(Paragraph(esc(block), styles["BodyJust"]))

        calculations = content.get("calculations") or {}
        if calculations:
            story.append(PageBreak())
            story.append(Paragraph("Cálculos e dimensionamentos", styles["H1Pt"]))
            for calc_name, calc_data in calculations.items():
                if "error" in calc_data:
                    continue
                story.append(
                    Paragraph(
                        f"<b>{esc(calc_name)}:</b> {esc(calc_data.get('value', ''))} "
                        f"{esc(calc_data.get('unit', ''))}<br/>"
                        f"<i>Fórmula: {esc(calc_data.get('formula', ''))}</i>",
                        styles["BodyJust"],
                    )
                )

        normas = content.get("normas") or []
        if normas:
            story.append(PageBreak())
            story.append(Paragraph("Referências normativas", styles["H1Pt"]))
            items = [ListItem(Paragraph(esc(str(norma)), styles["BodyJust"])) for norma in normas]
            story.append(ListFlowable(items, bulletType="bullet"))

        story.append(PageBreak())
        story.append(Paragraph("Dados do projeto e responsabilidade técnica", styles["H1Pt"]))
        table_data = [
            ["Obra", display(metadata.get("obra"))],
            ["Local", display(metadata.get("local") or shared.get("localizacao"))],
            ["Município", display(shared.get("municipio"))],
            ["UF", display(shared.get("uf"))],
            ["Responsável técnico", display(professional.get("full_name") or metadata.get("responsavel"))],
            ["CREA", display(metadata.get("crea"))],
            ["Empresa", display(professional.get("company_name"))],
            ["Data", datetime.now().strftime("%d/%m/%Y")],
        ]
        table = Table(table_data, colWidths=[5 * cm, 11 * cm])
        table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), "Times-Roman"),
                    ("FONTSIZE", (0, 0), (-1, -1), 10),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BACKGROUND", (0, 0), (0, -1), colors.Color(0.95, 0.95, 0.95)),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 0.6 * cm))
        story.append(
            Paragraph(
                "Declaro que as informações técnicas deste memorial foram elaboradas sob "
                "minha responsabilidade profissional, em conformidade com as normas vigentes.",
                styles["BodyJust"],
            )
        )

        def _footer(canvas, doc):
            canvas.saveState()
            canvas.setFont("Times-Roman", 9)
            canvas.drawCentredString(A4[0] / 2, 1.4 * cm, f"Página {doc.page}")
            canvas.restoreState()

        document.build(story, onFirstPage=_footer, onLaterPages=_footer)
        buffer.seek(0)
        return buffer
